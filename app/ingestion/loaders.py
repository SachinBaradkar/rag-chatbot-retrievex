"""
Document loaders — one function per file type, unified interface.
Returns list of {"text": str, "metadata": dict}.
"""

from __future__ import annotations
import csv
import json
import logging
import io
from pathlib import Path
from typing import List, Dict, Any

logger = logging.getLogger(__name__)

RawPage = Dict[str, Any]  # {"text": str, "metadata": dict}


# ── PDF ───────────────────────────────────────────────────────────────────────

def load_pdf(file_path: Path) -> List[RawPage]:
    try:
        import pdfplumber
        pages: List[RawPage] = []
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append({
                        "text": text,
                        "metadata": {
                            "source": file_path.name,
                            "page": i,
                            "file_type": "pdf",
                        },
                    })
        return pages
    except Exception as e:
        logger.error("PDF load error %s: %s", file_path, e)
        return []


# ── TXT / MD ──────────────────────────────────────────────────────────────────

def load_text(file_path: Path) -> List[RawPage]:
    try:
        text = file_path.read_text(encoding="utf-8", errors="replace")
        if not text.strip():
            return []
        return [{
            "text": text,
            "metadata": {
                "source": file_path.name,
                "page": 1,
                "file_type": file_path.suffix.lstrip("."),
            },
        }]
    except Exception as e:
        logger.error("Text load error %s: %s", file_path, e)
        return []


# ── DOCX ──────────────────────────────────────────────────────────────────────

def load_docx(file_path: Path) -> List[RawPage]:
    try:
        from docx import Document
        doc = Document(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)
        if not text.strip():
            return []
        return [{
            "text": text,
            "metadata": {
                "source": file_path.name,
                "page": 1,
                "file_type": "docx",
            },
        }]
    except Exception as e:
        logger.error("DOCX load error %s: %s", file_path, e)
        return []


# ── CSV ───────────────────────────────────────────────────────────────────────

def load_csv(file_path: Path) -> List[RawPage]:
    try:
        rows = []
        with open(file_path, newline="", encoding="utf-8", errors="replace") as f:
            reader = csv.DictReader(f)
            for row in reader:
                rows.append(", ".join(f"{k}: {v}" for k, v in row.items()))
        if not rows:
            return []
        # Group rows into chunks of 50 so retrieval is manageable
        chunk_size = 50
        pages: List[RawPage] = []
        for i in range(0, len(rows), chunk_size):
            chunk = rows[i:i + chunk_size]
            pages.append({
                "text": "\n".join(chunk),
                "metadata": {
                    "source": file_path.name,
                    "page": (i // chunk_size) + 1,
                    "file_type": "csv",
                },
            })
        return pages
    except Exception as e:
        logger.error("CSV load error %s: %s", file_path, e)
        return []


# ── XLSX ──────────────────────────────────────────────────────────────────────

def load_xlsx(file_path: Path) -> List[RawPage]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        pages: List[RawPage] = []
        for sheet in wb.worksheets:
            rows = []
            headers = None
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i == 0:
                    headers = [str(c) if c is not None else "" for c in row]
                    continue
                if headers:
                    line = ", ".join(f"{h}: {v}" for h, v in zip(headers, row) if v is not None)
                    if line.strip():
                        rows.append(line)
            if rows:
                pages.append({
                    "text": "\n".join(rows),
                    "metadata": {
                        "source": file_path.name,
                        "page": 1,
                        "sheet": sheet.title,
                        "file_type": "xlsx",
                    },
                })
        return pages
    except Exception as e:
        logger.error("XLSX load error %s: %s", file_path, e)
        return []


# ── JSON ──────────────────────────────────────────────────────────────────────

def load_json(file_path: Path) -> List[RawPage]:
    try:
        data = json.loads(file_path.read_text(encoding="utf-8"))
        text = json.dumps(data, indent=2)
        return [{
            "text": text,
            "metadata": {
                "source": file_path.name,
                "page": 1,
                "file_type": "json",
            },
        }]
    except Exception as e:
        logger.error("JSON load error %s: %s", file_path, e)
        return []


# ── Dispatcher ────────────────────────────────────────────────────────────────

_LOADERS = {
    ".pdf":  load_pdf,
    ".txt":  load_text,
    ".md":   load_text,
    ".docx": load_docx,
    ".csv":  load_csv,
    ".xlsx": load_xlsx,
    ".json": load_json,
}


def load_document(file_path: Path) -> List[RawPage]:
    ext = file_path.suffix.lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        logger.warning("Unsupported file type: %s", ext)
        return []
    return loader(file_path)
